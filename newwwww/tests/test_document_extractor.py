"""Tests for document extraction: PDF, DOCX, tables, corrupt/empty docs, and format verification.
"""
import io
import pytest
import pymupdf
from docx import Document

from app.exceptions import DocumentExtractionError, FileTooLargeError, InvalidFileFormatError
from app.services.document_extractor import DocumentExtractor
from app.utils.security import detect_file_type, validate_file_size


def make_synthetic_pdf(text: str) -> bytes:
    """Helper to generate a valid PDF byte string with extractable text."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def make_scanned_pdf() -> bytes:
    """Helper to generate a synthetic image-only/scanned PDF without text."""
    doc = pymupdf.open()
    temp_page = doc.new_page()
    shape = temp_page.new_shape()
    shape.draw_rect(pymupdf.Rect(50, 50, 150, 150))
    shape.finish(color=(1, 0, 0), fill=(0, 1, 0))
    shape.commit()
    pix = temp_page.get_pixmap()
    doc.close()

    scanned_doc = pymupdf.open()
    scanned_page = scanned_doc.new_page()
    scanned_page.insert_image(pymupdf.Rect(50, 50, 150, 150), pixmap=pix)
    pdf_bytes = scanned_doc.tobytes()
    scanned_doc.close()
    return pdf_bytes


def make_synthetic_docx(paragraphs: list, table_data: list = None) -> bytes:
    """Helper to generate a valid DOCX byte string with paragraphs and tables."""
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = str(val)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocumentExtractor:
    """Test suite for DocumentExtractor."""

    def test_valid_pdf_extraction(self):
        sample_text = "John Doe\nEducation: B.Tech in Computer Science\nTechnical Skills: Python, SQL"
        pdf_bytes = make_synthetic_pdf(sample_text)

        text, warnings = DocumentExtractor.extract_from_pdf(pdf_bytes)
        assert "John Doe" in text
        assert "Python" in text
        assert isinstance(warnings, list)

    def test_valid_docx_extraction(self):
        paragraphs = ["Jane Smith", "Technical Skills: Java, C++", "Projects: Distributed Cache"]
        docx_bytes = make_synthetic_docx(paragraphs)

        text, warnings = DocumentExtractor.extract_from_docx(docx_bytes)
        assert "Jane Smith" in text
        assert "Distributed Cache" in text

    def test_docx_table_extraction(self):
        paragraphs = ["Candidate Experience"]
        table_data = [
            ["Company", "Role", "Duration"],
            ["Acme Corp", "Software Intern", "Summer 2025"]
        ]
        docx_bytes = make_synthetic_docx(paragraphs, table_data=table_data)

        text, warnings = DocumentExtractor.extract_from_docx(docx_bytes)
        assert "Acme Corp | Software Intern | Summer 2025" in text

    def test_empty_pdf_document(self):
        doc = pymupdf.open()
        doc.new_page()  # Page with no text
        empty_pdf = doc.tobytes()
        doc.close()

        with pytest.raises(DocumentExtractionError) as exc_info:
            DocumentExtractor.extract_from_pdf(empty_pdf)
        assert "no extractable text" in str(exc_info.value).lower()

    def test_scanned_image_only_pdf_detection(self):
        scanned_bytes = make_scanned_pdf()
        with pytest.raises(DocumentExtractionError) as exc_info:
            DocumentExtractor.extract_from_pdf(scanned_bytes)
        assert "OCR is not supported" in str(exc_info.value)

    def test_corrupted_pdf_document(self):
        corrupted_bytes = b"%PDF-1.4\ncorrupted content that cannot be parsed %%EOF"
        with pytest.raises(DocumentExtractionError):
            DocumentExtractor.extract_from_pdf(corrupted_bytes)

    def test_empty_docx_document(self):
        doc = Document()
        buf = io.BytesIO()
        doc.save(buf)
        empty_docx = buf.getvalue()

        with pytest.raises(DocumentExtractionError) as exc_info:
            DocumentExtractor.extract_from_docx(empty_docx)
        assert "no readable text" in str(exc_info.value).lower()

    def test_unsupported_file_extension(self):
        with pytest.raises(InvalidFileFormatError) as exc_info:
            detect_file_type(b"some content", "resume.txt")
        assert "Unsupported file extension" in str(exc_info.value)

    def test_content_extension_mismatch(self):
        fake_pdf = b"This is plainly a text file and not a PDF header."
        with pytest.raises(InvalidFileFormatError) as exc_info:
            detect_file_type(fake_pdf, "resume.pdf")
        assert "%PDF-" in str(exc_info.value)

    def test_oversized_file_validation(self):
        # Limit to 1 MB for testing
        with pytest.raises(FileTooLargeError):
            validate_file_size(2 * 1024 * 1024, max_mb=1)

    def test_zero_byte_file(self):
        with pytest.raises(InvalidFileFormatError) as exc_info:
            detect_file_type(b"", "resume.pdf")
        assert "empty (0 bytes)" in str(exc_info.value)
