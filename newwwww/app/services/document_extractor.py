"""Document extraction service for extracting text from PDF and DOCX files.
Utilizes PyMuPDF for PDF documents and python-docx for DOCX documents.
"""
import io
import re
from typing import Dict, List, Tuple
import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.exceptions import DocumentExtractionError


class DocumentExtractor:
    """Handles text extraction from PDF and DOCX documents with structure preservation."""

    @staticmethod
    def extract_from_pdf(content: bytes) -> Tuple[str, List[str]]:
        """Extract plain text from PDF bytes using PyMuPDF.

        Returns (extracted_text, warnings).
        Raises DocumentExtractionError if the document is encrypted, corrupt,
        empty, or scanned/image-only without an extractable text layer.
        """
        warnings: List[str] = []
        try:
            doc = pymupdf.open(stream=content, filetype="pdf")
        except Exception as e:
            raise DocumentExtractionError(f"Failed to read PDF document: {str(e)}") from e

        try:
            if doc.is_encrypted:
                raise DocumentExtractionError("PDF is encrypted or password-protected and cannot be read.")

            page_count = len(doc)
            if page_count == 0:
                raise DocumentExtractionError("PDF contains 0 pages.")

            extracted_pages: List[str] = []
            total_images = 0

            for page_index in range(page_count):
                page = doc[page_index]
                text = page.get_text()
                extracted_pages.append(text.strip())
                images = page.get_images()
                total_images += len(images)

            full_text = "\n\n".join(p for p in extracted_pages if p)
            stripped_text = full_text.strip()

            # Detection for scanned / image-only PDFs
            # If text is extremely short or empty, but images exist, it's likely a scanned PDF
            if len(stripped_text) < 30 and total_images > 0:
                raise DocumentExtractionError(
                    "Scanned or image-only PDF detected. OCR is not supported by this parser. "
                    "Please provide a text-based PDF or DOCX resume."
                )

            if not stripped_text:
                raise DocumentExtractionError(
                    "PDF document contains no extractable text. The file appears to be blank or rendered as raster graphics."
                )

            return stripped_text, warnings

        finally:
            doc.close()

    @staticmethod
    def extract_from_docx(content: bytes) -> Tuple[str, List[str]]:
        """Extract text from DOCX bytes preserving paragraphs and table contents in order.

        Returns (extracted_text, warnings).
        Raises DocumentExtractionError if empty or unreadable.
        """
        warnings: List[str] = []
        try:
            doc = Document(io.BytesIO(content))
        except Exception as e:
            raise DocumentExtractionError(f"Failed to parse DOCX document: {str(e)}") from e

        ordered_blocks: List[str] = []

        # Iterate over document body elements in document order to preserve paragraph and table sequence
        try:
            for child in doc.element.body:
                tag = child.tag
                if tag.endswith("p"):
                    p = Paragraph(child, doc)
                    text = p.text.strip()
                    if text:
                        ordered_blocks.append(text)
                elif tag.endswith("tbl"):
                    table = Table(child, doc)
                    table_rows: List[str] = []
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            # Remove duplicate adjacent cell texts caused by merged cells
                            dedup_cells = []
                            for c in row_cells:
                                if not dedup_cells or dedup_cells[-1] != c:
                                    dedup_cells.append(c)
                            table_rows.append(" | ".join(dedup_cells))
                    if table_rows:
                        ordered_blocks.append("\n".join(table_rows))
        except Exception as e:
            # Fallback to standard paragraphs and tables iteration if element parsing encountered an issue
            warnings.append(f"Structured element traversal notice: {str(e)}")
            ordered_blocks.clear()
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    ordered_blocks.append(t)
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        ordered_blocks.append(" | ".join(row_cells))

        full_text = "\n\n".join(ordered_blocks).strip()
        if not full_text:
            raise DocumentExtractionError("DOCX document contains no readable text or is empty.")

        return full_text, warnings

    @classmethod
    def extract(cls, content: bytes, file_type: str) -> Tuple[str, List[str]]:
        """Dispatch document extraction based on verified file type ('pdf' or 'docx')."""
        if file_type == "pdf":
            return cls.extract_from_pdf(content)
        elif file_type == "docx":
            return cls.extract_from_docx(content)
        else:
            raise DocumentExtractionError(f"Unsupported file type '{file_type}'. Supported: 'pdf', 'docx'.")

