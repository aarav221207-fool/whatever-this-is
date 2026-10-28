import io
import re
from typing import Dict, Optional, Tuple

try:
    import pymupdf as fitz  # Modern PyMuPDF import
except ImportError:
    import fitz  # Legacy fallback
import docx  # python-docx

from app.exceptions import DocumentEmptyError, MalformedDocumentError, ScannedPDFError


def normalize_whitespace(text: str) -> str:
    """
    Normalizes whitespace while preserving readable paragraph boundaries.
    Replaces 3+ consecutive newlines with 2 newlines and removes redundant indentation.
    """
    if not text:
        return ""
    
    # Replace non-breaking spaces with normal spaces
    text = text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    
    # Normalize tabs to 2 spaces
    text = text.replace("\t", "  ")
    
    # Clean up trailing spaces per line
    lines = [line.strip() for line in text.split("\n")]
    
    # Reassemble and collapse more than 2 consecutive newlines
    reassembled = "\n".join(lines)
    normalized = re.sub(r"\n{3,}", "\n\n", reassembled)
    
    return normalized.strip()


class DocumentExtractor:
    """
    Extracts text and metadata from PDF and DOCX files.
    Detects scanned image PDFs, empty documents, corrupted files, and tables.
    """

    @classmethod
    def extract_from_bytes(cls, content: bytes, extension: str) -> Tuple[str, Dict[str, Optional[str]]]:
        """
        Extracts normalized text and metadata from raw file bytes.
        Returns:
            Tuple of (extracted_text, metadata_dict)
        Raises:
            DocumentEmptyError, ScannedPDFError, MalformedDocumentError
        """
        ext = extension.lower()
        if not ext.startswith("."):
            ext = f".{ext}"

        if ext == ".pdf":
            return cls._extract_pdf(content)
        elif ext == ".docx":
            return cls._extract_docx(content)
        else:
            raise MalformedDocumentError(f"Unsupported document format: {ext}")

    @classmethod
    def _extract_pdf(cls, content: bytes) -> Tuple[str, Dict[str, Optional[str]]]:
        """
        Extract text and metadata from PDF using PyMuPDF (fitz).
        Detects scanned/image-only PDFs.
        """
        try:
            doc = fitz.open(stream=content, filetype="pdf")
        except Exception as e:
            raise MalformedDocumentError(f"Corrupted or invalid PDF structure: {str(e)}")

        try:
            total_pages = len(doc)
            if total_pages == 0:
                raise DocumentEmptyError("PDF contains 0 pages.")

            extracted_chunks = []
            total_chars = 0
            has_images = False

            for page_idx in range(total_pages):
                page = doc[page_idx]
                page_text = page.get_text("text") or ""
                trimmed = page_text.strip()
                total_chars += len(trimmed)
                
                if trimmed:
                    extracted_chunks.append(f"--- Page {page_idx + 1} ---\n{trimmed}")

                # Check if page contains images
                images = page.get_images()
                if images and len(images) > 0:
                    has_images = True

            # Detect likely scanned PDF:
            # If document has pages but text is negligible (e.g. fewer than 35 characters for the whole document)
            if total_chars < 35:
                if has_images or total_chars == 0:
                    raise ScannedPDFError(
                        "Document appears to be a scanned image-only PDF with no extractable text layer. "
                        "OCR (Optical Character Recognition) is currently unsupported. "
                        "Please upload a text-based PDF or DOCX file."
                    )
                else:
                    raise DocumentEmptyError("PDF contains no readable text content.")

            raw_text = "\n\n".join(extracted_chunks)
            normalized_text = normalize_whitespace(raw_text)

            if not normalized_text:
                raise DocumentEmptyError("Extracted text is empty after whitespace normalization.")

            metadata = {
                "title": doc.metadata.get("title") or None,
                "author": doc.metadata.get("author") or None,
                "subject": doc.metadata.get("subject") or None,
                "pages": str(total_pages)
            }

            return normalized_text, metadata

        finally:
            doc.close()

    @classmethod
    def _extract_docx(cls, content: bytes) -> Tuple[str, Dict[str, Optional[str]]]:
        """
        Extract text, tables, and core properties from DOCX using python-docx.
        """
        try:
            doc_stream = io.BytesIO(content)
            doc = docx.Document(doc_stream)
        except Exception as e:
            raise MalformedDocumentError(f"Corrupted or invalid DOCX document: {str(e)}")

        chunks = []

        # 1. Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                chunks.append(text)

        # 2. Extract tables
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate identical adjacent merged cells
                deduped = []
                for c in row_cells:
                    if not deduped or deduped[-1] != c:
                        deduped.append(c)
                if deduped:
                    table_lines.append(" | ".join(deduped))
            if table_lines:
                chunks.append("\n".join(table_lines))

        raw_text = "\n\n".join(chunks)
        normalized_text = normalize_whitespace(raw_text)

        if not normalized_text:
            raise DocumentEmptyError("DOCX document contains no readable paragraphs or table text.")

        metadata = {}
        try:
            props = doc.core_properties
            metadata["title"] = props.title or None
            metadata["author"] = props.author or None
            metadata["subject"] = props.subject or None
        except Exception:
            pass

        return normalized_text, metadata
