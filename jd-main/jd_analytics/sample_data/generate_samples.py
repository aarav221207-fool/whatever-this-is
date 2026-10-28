"""
Script to generate synthetic sample Job Descriptions (PDF and DOCX) for testing and development.
Run:
    python sample_data/generate_samples.py
"""
import os
import sys

# Ensure parent directory is in path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

try:
    import pymupdf as fitz
except ImportError:
    import fitz
import docx  # python-docx

SAMPLE_DIR = os.path.dirname(os.path.abspath(__file__))


def generate_sample_pdf():
    pdf_path = os.path.join(SAMPLE_DIR, "backend_engineer_jd.pdf")
    doc = fitz.open()

    p1 = doc.new_page()
    text_p1 = (
        "Company: NovaScale Technologies\n"
        "Job Title: Senior Backend Engineer\n\n"
        "About NovaScale Technologies:\n"
        "NovaScale builds mission-critical high-throughput microservices for modern enterprise systems.\n\n"
        "Responsibilities:\n"
        "- Design and implement scalable services that communicate over HTTP using REST APIs and gRPC.\n"
        "- Work with relational data and write complex queries against PostgreSQL databases.\n"
        "- Investigate algorithmic complexity and optimize critical database queries for high concurrency.\n"
        "- Build and maintain CI/CD pipelines and automated testing suites.\n\n"
        "Minimum Qualifications (Required):\n"
        "- 4+ years of professional backend software development experience.\n"
        "- High proficiency with Python or Go.\n"
        "- Strong understanding of Data Structures & Algorithms and Big-O computational complexity.\n"
        "- Solid grounding in Object-Oriented Design and design patterns.\n"
        "- Experience with relational databases, SQL indexing, and schema design.\n"
        "- Experience with Docker containerization and AWS infrastructure.\n"
        "- Strong problem solving aptitude and collaborative team communication."
    )
    p1.insert_text((50, 72), text_p1, fontsize=10)

    p2 = doc.new_page()
    text_p2 = (
        "Preferred Qualifications (Bonus):\n"
        "- Hands-on experience with Kubernetes orchestration.\n"
        "- Experience training machine learning models or building predictive pipelines.\n"
        "- Familiarity with Linux kernel concepts, process memory, and multithreading.\n"
        "- Experience with Redis caching and Apache Kafka event-driven architectures."
    )
    p2.insert_text((50, 72), text_p2, fontsize=10)

    doc.save(pdf_path)
    doc.close()
    print(f"Generated sample PDF at: {pdf_path}")


def generate_sample_docx():
    docx_path = os.path.join(SAMPLE_DIR, "cloud_architect_jd.docx")
    doc = docx.Document()

    doc.add_heading("Apex Cloud Networks", level=1)
    doc.add_heading("Role: Principal Cloud & Distributed Systems Architect", level=2)

    doc.add_paragraph(
        "Apex Cloud Networks is looking for an experienced systems architect to lead our next-generation cloud foundation."
    )

    doc.add_heading("Required Qualifications:", level=3)
    doc.add_paragraph("• 8+ years designing large-scale distributed systems and high-availability cloud solutions.")
    doc.add_paragraph("• Expert knowledge of AWS, Google Cloud Platform (GCP), and Terraform.")
    doc.add_paragraph("• Deep understanding of computer networking, TCP/IP, load balancing, DNS, and TLS certificates.")
    doc.add_paragraph("• Strong proficiency in Go, Python, or Rust for systems infrastructure.")
    doc.add_paragraph("• Proven track record leading technical reviews and cross-functional engineering communication.")

    doc.add_heading("Preferred Skills:", level=3)
    doc.add_paragraph("• Experience building automated CI/CD deployments with GitHub Actions.")
    doc.add_paragraph("• Familiarity with AI data processing pipelines and PyTorch inference servers.")

    # Table with structured requirements
    table = doc.add_table(rows=4, cols=3)
    table_data = [
        ["Domain", "Required Technology", "Experience"],
        ["Cloud", "AWS, Terraform, Kubernetes", "5+ Years"],
        ["Networking", "HTTP/HTTPS, gRPC, WebSockets", "5+ Years"],
        ["Databases", "PostgreSQL, Redis", "4+ Years"]
    ]
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    doc.save(docx_path)
    print(f"Generated sample DOCX at: {docx_path}")


if __name__ == "__main__":
    generate_sample_pdf()
    generate_sample_docx()
