import io
import pytest
try:
    import pymupdf as fitz
except ImportError:
    import fitz
import docx  # python-docx


@pytest.fixture
def make_pdf():
    """
    Factory fixture to create text-based synthetic PDF files in memory.
    Supports multi-page generation.
    """
    def _create(pages_text: list[str]) -> bytes:
        doc = fitz.open()
        for text in pages_text:
            page = doc.new_page()
            # Insert text at top-left
            page.insert_text((50, 72), text, fontsize=11)
        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes
    return _create


@pytest.fixture
def make_docx():
    """
    Factory fixture to create synthetic DOCX files in memory with paragraphs and tables.
    """
    def _create(paragraphs: list[str], tables_data: list[list[list[str]]] = None) -> bytes:
        doc = docx.Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        
        if tables_data:
            for t_data in tables_data:
                if t_data and len(t_data) > 0:
                    rows = len(t_data)
                    cols = len(t_data[0])
                    table = doc.add_table(rows=rows, cols=cols)
                    for r_idx, row in enumerate(t_data):
                        for c_idx, val in enumerate(row):
                            table.cell(r_idx, c_idx).text = val
        
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
    return _create


@pytest.fixture
def make_scanned_pdf():
    """
    Creates an image-only synthetic PDF with no extractable text layer to simulate scanned resumes/JDs.
    """
    def _create() -> bytes:
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        # Create a small pixmap drawing (simulating a scanned page)
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), False)
        pix.clear_with(255)
        # Insert image into page
        page.insert_image(fitz.Rect(50, 50, 400, 500), pixmap=pix)
        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes
    return _create


@pytest.fixture
def sample_backend_jd_text():
    return """
Company: Acme Cloud Solutions
Job Title: Senior Backend Developer

About Us:
Acme Cloud Solutions builds enterprise distributed systems.

Requirements:
- 5+ years of software engineering experience.
- Strong proficiency in Python or Java for backend services.
- Experience with SQL and relational databases like PostgreSQL.
- Must have experience writing complex queries and relational data modeling.
- Solid understanding of Data Structures & Algorithms and algorithmic complexity.
- Experience designing scalable services that communicate over HTTP using REST APIs.
- Experience with Docker and AWS cloud deployments.
- Strong problem-solving aptitude and cross-functional communication skills.

Preferred Qualifications:
- Familiarity with Kubernetes and microservices architecture.
- Bonus: Hands-on experience with Machine Learning and building predictive models.
- Familiarity with Linux shell scripting and multithreading.
"""
